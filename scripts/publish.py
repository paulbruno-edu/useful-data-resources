"""Check the workbook, update the website's data, and publish it. One step.

    python scripts/publish.py              check, export, commit, and push
    python scripts/publish.py --dry-run    check and export only; nothing is committed or pushed

Close the workbook in Excel first. The steps:
  1. Check the workbook: every source has a name and a link, the dropdown columns use
     values from the lists tab, and no name appears twice. Problems stop the run.
  2. Export it to the website (export_site.py): the data/ folder, the plain-HTML lists in
     index.html, and sitemap.xml.
  3. Commit those files and push them to GitHub. GitHub Pages republishes in about a minute.
  4. List the five sources reviewed least recently, as a reminder to check a few by hand.

The workbook's location is set in export_site.py (WORKBOOK); this repository is the website.
"""
import datetime
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parent))
import export_site  # noqa: E402

BOOK = export_site.BOOK
REPO = export_site.SITE_REPO
# What the export writes, and so what this script commits
GENERATED = ["data", "index.html", "sitemap.xml"]

# Columns whose values must come from the lists tab, and columns that may hold several
# values separated by semicolons
CHECKED = {"Geography": "Geography", "Level": "Level", "Unit of analysis": "Unit of analysis",
           "Access": "Access", "Kind": "Kind", "Link status": "Link status", "Topics": "Topics",
           "Start here": "Start here"}
MULTI = {"Level", "Topics", "Unit of analysis"}


def check_workbook():
    """Return a list of problems; empty means the workbook is ready to publish."""
    wb = load_workbook(BOOK, read_only=True)
    lists = {}
    lt = wb["lists"]
    rows = list(lt.iter_rows(values_only=True))
    for i, name in enumerate(rows[0]):
        if name:
            lists[name] = {r[i] for r in rows[1:] if r[i]}

    problems, seen = [], {}
    for tab in ("resources", "learning_resources"):
        data = list(wb[tab].iter_rows(values_only=True))
        hdr = {v: i for i, v in enumerate(data[0])}
        for n, row in enumerate(data[1:], 2):
            if not any(row):
                continue
            name = row[hdr["Name"]]
            where = f"{tab} row {n}" + (f" ({name})" if name else "")
            if not name:
                problems.append(f"{where}: no name")
                continue
            if name in seen:
                problems.append(f"{where}: same name as {seen[name]}")
            seen[name] = f"{tab} row {n}"
            link = row[hdr["Link"]] or ""
            if not str(link).startswith("http"):
                problems.append(f"{where}: link is missing or doesn't start with http")
            if tab != "resources":
                continue
            for col, list_name in CHECKED.items():
                val = row[hdr[col]]
                if not val:
                    if col not in ("Start here",):
                        problems.append(f"{where}: {col} is blank")
                    continue
                vals = [v.strip() for v in str(val).split(";")] if col in MULTI else [str(val).strip()]
                for v in vals:
                    if v not in lists.get(list_name, set()):
                        problems.append(f"{where}: {col} \"{v}\" isn't in the lists tab")

    # Every Start here source needs a row in the start_here tab, and every row there must
    # name a Start here source and one of the listed groups
    res = list(wb["resources"].iter_rows(values_only=True))
    rh = {v: i for i, v in enumerate(res[0])}
    starts = {r[rh["Name"]] for r in res[1:] if r[rh["Name"]] and str(r[rh["Start here"]] or "").lower() == "yes"}
    sh = list(wb["start_here"].iter_rows(values_only=True))
    shh = {v: i for i, v in enumerate(sh[0])}
    listed = set()
    for n, row in enumerate(sh[1:], 2):
        name = row[shh["Name"]]
        if not name:
            continue
        listed.add(name)
        if name not in starts:
            problems.append(f"start_here row {n} ({name}): not marked Start here in the resources tab (or the names differ)")
        if row[shh["Group"]] not in lists.get("Start here group", set()):
            problems.append(f"start_here row {n} ({name}): Group \"{row[shh['Group']]}\" isn't in the lists tab")
    for name in sorted(starts - listed):
        problems.append(f"{name}: marked Start here but has no row in the start_here tab")
    return problems


def review_reminder(n=5):
    """The n sources reviewed least recently (never-reviewed first, Start here sources first)."""
    wb = load_workbook(BOOK, read_only=True)
    data = list(wb["resources"].iter_rows(values_only=True))
    hdr = {v: i for i, v in enumerate(data[0])}
    items = []
    for row in data[1:]:
        if not row[hdr["Name"]]:
            continue
        reviewed = row[hdr["Last reviewed"]]
        if isinstance(reviewed, datetime.datetime):
            reviewed = reviewed.date().isoformat()
        start = str(row[hdr["Start here"]] or "").lower() == "yes"
        items.append((reviewed or "", not start, row[hdr["Name"]], row[hdr["Link"]]))
    items.sort()
    return items[:n]


def git(*args):
    return subprocess.run(["git", *args], cwd=REPO, capture_output=True, text=True)


def main():
    dry = "--dry-run" in sys.argv
    lock = BOOK.with_name("~$" + BOOK.name)
    if lock.exists():
        sys.exit("The workbook is open in Excel. Save and close it, then run this again.")

    print("Checking the workbook...")
    problems = check_workbook()
    if problems:
        print(f"\nFound {len(problems)} problem(s). Fix these in the workbook, then run this again:\n")
        for p in problems:
            print("  - " + p)
        sys.exit(1)
    print("  OK")

    if not dry:
        # Pick up anything changed on GitHub first (GitHub sometimes commits there itself)
        r = git("pull", "--ff-only", "-q")
        if r.returncode:
            sys.exit("Couldn't update the website folder from GitHub:\n" + r.stderr)

    print("Exporting...")
    export_site.main()

    changed = git("status", "--porcelain", "--", *GENERATED).stdout.strip()
    other = [l for l in git("status", "--porcelain").stdout.splitlines()
             if not any(l[3:] == g or l[3:].startswith(g + "/") for g in GENERATED)]
    if not changed:
        print("\nThe website already matches the workbook. Nothing to publish.")
    elif dry:
        print("\nDry run: the website's data files are updated but not published.")
    else:
        today = datetime.date.today().isoformat()
        git("add", "--", *GENERATED)
        r = git("commit", "-q", "-m", f"Update the list ({today})")
        if r.returncode:
            sys.exit("Commit failed:\n" + r.stdout + r.stderr)
        r = git("push", "-q")
        if r.returncode:
            sys.exit("Push failed (the change is committed locally; push it with GitHub Desktop):\n" + r.stderr)
        print("\nPublished. https://data.paul-bruno.com updates in about a minute.")
    if other:
        print("\nNote: the website folder also has changes this script doesn't publish:")
        for l in other:
            print("  " + l)

    print("\nDue for a review (check each source's contents and access, then set Last reviewed):")
    for reviewed, _, name, link in review_reminder():
        print(f"  {reviewed or 'never':10}  {name}\n{'':14}{link}")


if __name__ == "__main__":
    main()
