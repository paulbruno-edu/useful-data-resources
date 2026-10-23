"""Export useful_data_resources.xlsx to the data files the website reads.

    python scripts/export_site.py

Usually you'll run publish.py instead, which checks the workbook, runs this export, and
publishes the result. Run this alone to update the website's files without publishing them.

The workbook lives outside this repository (in Google Drive); WORKBOOK below says where. To use a
different copy without editing this file, set the USEFUL_DATA_WORKBOOK environment variable to
its path.

Writes, under the repository's data/ folder:
  resources.js    the data the page loads (works when opened from disk, too)
  resources.json  same data, for anyone who wants it
  resources.csv   the resources tab, for download
and, in the repository itself, the plain-HTML copy of the lists inside index.html (for search
engines) and sitemap.xml.

Review notes stay in the workbook; they aren't published.
"""
import csv
import datetime
import hashlib
import html
import json
import os
import re
from pathlib import Path

from openpyxl import load_workbook

# Where the workbook lives. Change this line if it moves.
WORKBOOK = r"C:\GoogleDrive\UIUC\Projects\usefuldataresources\useful_data_resources.xlsx"
BOOK = Path(os.environ.get("USEFUL_DATA_WORKBOOK", WORKBOOK))

SITE_REPO = Path(__file__).resolve().parent.parent  # this repository
OUT = SITE_REPO / "data"

RESOURCE_COLS = {
    "Name": "name", "Link": "link", "Description": "description", "Summary": "summary",
    "Geography": "geography", "State or country": "place", "Level": "level", "Topics": "topics", "Unit of analysis": "unit",
    "Access": "access", "Kind": "kind", "Start here": "start", "Link status": "status",
    "Last checked": "checked", "Added": "added",
}
LEARNING_COLS = {
    "Name": "name", "Link": "link", "Description": "description", "Type": "type",
    "Link status": "status", "Last checked": "checked",
}
START_COLS = {
    "Name": "name", "Group": "group", "Useful for": "useful", "Each row is": "row",
    "Getting started": "getting_started", "Years": "years", "Structure": "structure",
    "Documentation": "docs",
}


def read_tab(ws, cols):
    hdr = {c.value: c.column for c in ws[1]}
    rows = []
    for r in range(2, ws.max_row + 1):
        if not ws.cell(r, hdr["Name"]).value:
            continue
        row = {}
        for col, key in cols.items():
            v = ws.cell(r, hdr[col]).value
            if isinstance(v, datetime.datetime):  # a date typed in Excel
                v = v.date().isoformat()
            row[key] = "" if v is None else str(v).strip()
        rows.append(row)
    return rows


def main():
    wb = load_workbook(BOOK, read_only=True)
    resources = read_tab(wb["resources"], RESOURCE_COLS)
    learning = read_tab(wb["learning_resources"], LEARNING_COLS)
    start_here = read_tab(wb["start_here"], START_COLS)  # in the tab's order, which the site keeps

    for r in resources:
        r["topics"] = [t.strip() for t in r["topics"].split(";") if t.strip()]
        r["levels"] = [t.strip() for t in r["level"].split(";") if t.strip()]
        r["units"] = [t.strip() for t in r["unit"].split(";") if t.strip()]
        r["start"] = r["start"].lower() == "yes"
    checked = max(r["checked"] for r in resources + learning if r["checked"])

    data = {"updated": datetime.date.today().isoformat(), "checked": checked,
            "resources": resources, "start_here": start_here, "learning": learning}
    OUT.mkdir(parents=True, exist_ok=True)
    payload = json.dumps(data, ensure_ascii=False, indent=1)
    (OUT / "resources.json").write_text(payload, encoding="utf-8")
    (OUT / "resources.js").write_text(f"window.UDR = {payload};\n", encoding="utf-8")

    with open(OUT / "resources.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        cols = [c for c in RESOURCE_COLS if c != "Link status"]
        w.writerow(cols)
        for r in resources:
            vals = []
            for c in cols:
                v = r[RESOURCE_COLS[c]]
                vals.append("; ".join(v) if isinstance(v, list) else ("Yes" if v is True else "" if v is False else v))
            w.writerow(vals)

    # Version the data file by its contents, so browsers fetch new data as soon as it's
    # published instead of reusing a cached copy (GitHub Pages lets them keep it 10 minutes)
    version = hashlib.sha1(payload.encode("utf-8")).hexdigest()[:10]
    prerender(data, version)
    print(f"{len(resources)} resources, {len(learning)} learning resources, last checked {checked}")


# The page builds its lists with JavaScript, which search engines may not run. So the export
# also writes the lists into index.html as plain HTML (between the prerender markers), plus
# schema.org markup describing the catalog, and a sitemap. The script replaces the plain lists
# with the interactive ones when the page loads.
SITE_URL = "https://data.paul-bruno.com/"
SORT = lambda r: r["name"].lower().removeprefix("the ")  # noqa: E731  (same order as the page)


def fill(page, marker, content):
    start, end = f"<!-- prerender:{marker} -->", f"<!-- /prerender:{marker} -->"
    a, b = page.index(start) + len(start), page.index(end)
    return page[:a] + content + page[b:]


def prerender(data, version):
    e = html.escape
    live = [r for r in data["resources"] if r["status"] != "Broken"]
    cards = []
    for r in sorted(live, key=SORT):
        place = r["place"] if r["place"] and r["geography"] in ("US state", "US local", "International") else r["geography"]
        meta = "".join(f"<li>{e(v)}</li>" for v in [r["access"], ", ".join(r["levels"]), place, r["kind"]] if v)
        cards.append(f'<li class="result"><h3><a href="{e(r["link"])}" rel="noopener">{e(r["name"])}</a></h3>'
                     f'<p class="desc">{e(r["summary"] or r["description"])}</p><ul class="meta">{meta}</ul></li>')
    learning = "".join(f'<li><a href="{e(r["link"])}" rel="noopener">{e(r["name"])}</a><span>{e(r["description"])}</span></li>'
                       for r in sorted(data["learning"], key=SORT) if r["status"] != "Broken")

    catalog = {
        "@context": "https://schema.org",
        "@type": "DataCatalog",
        "name": "Useful Data Resources",
        "url": SITE_URL,
        "description": "Free and low-cost data sources for quantitative research in education, collected for doctoral students.",
        "creator": {"@type": "Person", "name": "Paul Bruno", "url": "https://www.paul-bruno.com/"},
        "dateModified": data["updated"],
        "inLanguage": "en",
        "dataset": [{"@type": "Dataset", "name": r["name"], "url": r["link"],
                     "description": r["summary"] or r["description"] or r["name"]}
                    for r in sorted(live, key=SORT)],
    }
    jsonld = '<script type="application/ld+json">' + json.dumps(catalog, ensure_ascii=False).replace("</", "<\\/") + "</script>"

    index = SITE_REPO / "index.html"
    page = index.read_text(encoding="utf-8")
    page = fill(page, "results", "".join(cards))
    page = fill(page, "learning", f"<ul>{learning}</ul>")
    page = fill(page, "jsonld", jsonld)
    page, n = re.subn(r'src="data/resources\.js\?v=[^"]*"', f'src="data/resources.js?v={version}"', page)
    assert n == 1, "index.html should load data/resources.js exactly once"
    index.write_text(page, encoding="utf-8")

    (SITE_REPO / "sitemap.xml").write_text(
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">\n'
        f"  <url><loc>{SITE_URL}</loc><lastmod>{data['updated']}</lastmod></url>\n"
        f"  <url><loc>{SITE_URL}guide.html</loc></url>\n"
        "</urlset>\n", encoding="utf-8")


if __name__ == "__main__":
    main()
