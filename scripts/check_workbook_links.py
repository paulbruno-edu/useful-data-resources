"""Recheck every link in the workbook (useful_data_resources.xlsx).

Updates "Link status" and "Last checked" on the resources and learning_resources
tabs, and prints anything that isn't OK. Close the workbook in Excel first.

    python scripts/check_workbook_links.py              # check and save
    python scripts/check_workbook_links.py --dry-run    # report only, don't save

Status rules:
  OK            200 at the same address (ignoring http->https, www, trailing slash)
  Moved         200 at a different address; the new address is printed so you can
                decide whether to update the link. Existing "Moved (link updated)"
                rows that now resolve cleanly become OK.
  Broken        404 or 410
  Check by hand anything else (403, timeouts): usually a site blocking scripts
"""
import concurrent.futures as cf
import datetime
import sys
import urllib.error
import urllib.request
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parent))
from export_site import BOOK  # noqa: E402  (the workbook's location is set there)
TABS = ["resources", "learning_resources"]
# Hosts that some sites redirect scripts to instead of serving the page
BOT_WALLS = ["perfdrive.com", "shieldsquare", "captcha"]
UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/130 Safari/537.36"


def norm(u):
    u = u.lower().split("#")[0].rstrip("/")
    for p in ("https://", "http://"):
        u = u.removeprefix(p)
    return u.removeprefix("www.")


def check(url):
    req = urllib.request.Request(url, headers={"User-Agent": UA})
    try:
        with urllib.request.urlopen(req, timeout=25) as r:
            final = r.geturl()
            if any(w in final for w in BOT_WALLS):
                return "Check by hand", "bot check page"
            return ("OK", final) if norm(final) == norm(url) else ("Moved (find new page)", final)
    except urllib.error.HTTPError as e:
        return ("Broken" if e.code in (404, 410) else "Check by hand", f"HTTP {e.code}")
    except Exception as e:
        return ("Check by hand", type(e).__name__)


def main():
    dry = "--dry-run" in sys.argv
    wb = load_workbook(BOOK)
    today = datetime.date.today().isoformat()
    jobs = []
    for tab in TABS:
        ws = wb[tab]
        hdr = {c.value: c.column for c in ws[1]}
        for row in range(2, ws.max_row + 1):
            url = ws.cell(row, hdr["Link"]).value
            if url and str(url).startswith("http"):
                jobs.append((ws, hdr, row, str(url)))

    with cf.ThreadPoolExecutor(16) as ex:
        results = list(ex.map(lambda j: check(j[3]), jobs))

    for (ws, hdr, row, url), (status, detail) in zip(jobs, results):
        name = ws.cell(row, hdr["Name"]).value
        ws.cell(row, hdr["Link status"]).value = status
        ws.cell(row, hdr["Last checked"]).value = today
        if status != "OK":
            print(f"{status:22} {ws.title}: {name}\n{'':22} {url} -> {detail}")

    if not dry:
        wb.save(BOOK)
        print(f"Saved {BOOK.name}")


if __name__ == "__main__":
    main()
